import io

from pptx import Presentation

from app.services.exports.deck import build_pitch_deck


def _report() -> dict:
    return {
        "scope": "site",
        "target_url": "https://acme.com",
        "overall_score": 62,
        "status": "warn",
        "layer_scores": {"SEO": 70, "AEO": 55, "GEO": 40},
        "executive_summary": "Acme leads on SEO foundations but trails on GEO citations.",
        "modules": [
            {"key": "technical_seo", "title": "Technical SEO", "layer": "SEO", "score": 72, "status": "warn"},
            {"key": "on_page", "title": "On-Page SEO", "layer": "SEO", "score": 68, "status": "warn"},
            {"key": "aeo_audit", "title": "AEO Audit", "layer": "AEO", "score": 55, "status": "warn"},
            {"key": "geo_audit", "title": "GEO Audit", "layer": "GEO", "score": 40, "status": "fail"},
        ],
        "roadmap": [
            {"phase": "SEO", "order": 1, "priority": "high", "layer": "SEO",
             "module": "technical_seo", "action": "Fix noindex", "how_to": "x",
             "effort": "low", "impact": "high", "rationale": None},
            {"phase": "AEO", "order": 2, "priority": "med", "layer": "AEO",
             "module": "aeo_audit", "action": "Add FAQ schema", "how_to": "x",
             "effort": "med", "impact": "high", "rationale": None},
            {"phase": "GEO", "order": 3, "priority": "high", "layer": "GEO",
             "module": "geo_audit", "action": "Earn citations", "how_to": "x",
             "effort": "high", "impact": "high", "rationale": None},
        ],
        "benchmarks": [{"metric": "title_length_max", "value": 60, "unit": "chars",
                        "source": "Moz (2024)", "source_url": "https://moz.com"}],
        "competitor_summary": [
            {"competitor": "Rival", "signal": "domain_rating", "us": 40, "them": 65, "gap": -25},
        ],
        "module_results": [
            {"module": "technical_seo", "score": 72, "status": "warn",
             "findings": [{"signal": "https", "value": 0.8, "benchmark": 1.0}],
             "recommendations": [{"action": "Serve all pages over HTTPS"}]},
            {"module": "aeo_audit", "score": 55, "status": "warn",
             "findings": [{"signal": "ai_overview_readiness", "value": 0.4, "benchmark": 1.0}],
             "recommendations": [{"action": "Add lead answer paragraphs"}]},
            {"module": "geo_audit", "score": 40, "status": "fail",
             "findings": [{"signal": "chatgpt_citation_rate", "value": 0.3, "benchmark": 0.4}],
             "recommendations": [{"action": "Increase citation-worthiness"}]},
        ],
        "synthesis_source": "deterministic",
    }


def _open(data: bytes) -> Presentation:
    return Presentation(io.BytesIO(data))


def test_deck_builds_with_slide_count_in_range():
    data = build_pitch_deck(_report(), {"name": "Acme", "domain": "acme.com"},
                            fetch_logo=lambda d: None)
    assert data[:2] == b"PK"  # pptx is a zip
    prs = _open(data)
    assert 25 <= len(list(prs.slides)) <= 30


def test_deck_has_key_takeaways_and_charts():
    data = build_pitch_deck(_report(), {"name": "Acme", "domain": "acme.com"},
                            fetch_logo=lambda d: None)
    prs = _open(data)
    takeaways = 0
    charts = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and "KEY TAKEAWAY" in shape.text_frame.text:
                takeaways += 1
            if shape.has_chart:
                charts += 1
    assert takeaways >= 12  # one per content slide
    assert charts >= 3  # scorecard, modules, GEO/AEO etc.


def test_deck_embeds_logo_when_provided():
    png = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01"
        b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    data = build_pitch_deck(_report(), {"name": "Acme", "domain": "acme.com"},
                            fetch_logo=lambda d: png)
    prs = _open(data)
    pics = sum(1 for slide in prs.slides for shape in slide.shapes if shape.shape_type == 13)
    assert pics >= 1  # logo embedded on the title slide


def test_deck_degrades_without_logo_or_competitors():
    rep = _report()
    rep["competitor_summary"] = []
    data = build_pitch_deck(rep, {"name": "Acme", "domain": ""}, fetch_logo=lambda d: None)
    assert data[:2] == b"PK"
    assert 25 <= len(list(_open(data).slides)) <= 30
