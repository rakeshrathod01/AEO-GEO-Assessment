"""Master Excel workbook for the Leadership Dashboard.

Executive Summary tab + all sections (module scores, prioritized roadmap, all
findings, all recommendations, competitor delta, benchmark sources).
"""

from __future__ import annotations

import io

from openpyxl import Workbook

from app.services.exports.excel import _write_sheet


def build_master_excel(report: dict) -> bytes:
    wb = Workbook()

    summary = wb.active
    summary.title = "Executive Summary"
    ls = report.get("layer_scores", {})
    _write_sheet(
        summary,
        ["Field", "Value"],
        [
            ["Scope", report.get("scope")],
            ["Target", report.get("target_url")],
            ["Overall score", report.get("overall_score")],
            ["Status", report.get("status")],
            ["SEO score", ls.get("SEO")],
            ["AEO score", ls.get("AEO")],
            ["GEO score", ls.get("GEO")],
            ["Synthesis", report.get("synthesis_source")],
            ["Summary", report.get("executive_summary")],
        ],
        [20, 100],
    )

    _write_sheet(
        wb.create_sheet("Module Scores"),
        ["Module", "Layer", "Score", "Status"],
        [[m["title"], m["layer"], m["score"], m["status"]] for m in report.get("modules", [])],
        [30, 10, 10, 10],
    )

    _write_sheet(
        wb.create_sheet("Roadmap"),
        ["#", "Phase", "Priority", "Module", "Action", "How To", "Effort", "Impact", "Rationale"],
        [
            [i["order"], i["phase"], i["priority"], i["module"], i["action"],
             i.get("how_to"), i.get("effort"), i.get("impact"), i.get("rationale")]
            for i in report.get("roadmap", [])
        ],
        [5, 8, 9, 18, 36, 44, 9, 9, 40],
    )

    # All findings / recommendations across modules.
    findings_rows, rec_rows = [], []
    for res in report.get("module_results", []):
        mod = res.get("module")
        for f in res.get("findings", []):
            findings_rows.append([mod, f.get("signal"), f.get("status"), f.get("value"),
                                  f.get("benchmark"), f.get("source"), f.get("evidence")])
        for r in res.get("recommendations", []):
            rec_rows.append([mod, r.get("priority"), r.get("layer"), r.get("action"),
                             r.get("how_to"), r.get("effort"), r.get("impact")])
    _write_sheet(
        wb.create_sheet("All Findings"),
        ["Module", "Signal", "Status", "Value", "Benchmark", "Source", "Evidence"],
        findings_rows, [18, 24, 9, 12, 12, 44, 40],
    )
    _write_sheet(
        wb.create_sheet("All Recommendations"),
        ["Module", "Priority", "Layer", "Action", "How To", "Effort", "Impact"],
        rec_rows, [18, 9, 8, 36, 56, 9, 9],
    )

    _write_sheet(
        wb.create_sheet("Competitor Delta"),
        ["Competitor", "Signal", "Us", "Them", "Gap"],
        [[d["competitor"], d["signal"], d["us"], d["them"], d["gap"]]
         for d in report.get("competitor_summary", [])],
        [22, 28, 12, 12, 12],
    )

    _write_sheet(
        wb.create_sheet("Benchmarks"),
        ["Metric", "Value", "Unit", "Source"],
        [[b["metric"], b.get("value"), b.get("unit"), b.get("source")]
         for b in report.get("benchmarks", [])],
        [28, 12, 10, 70],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
