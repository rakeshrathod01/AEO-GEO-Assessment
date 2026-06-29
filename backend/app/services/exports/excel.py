"""Excel export for a module analysis result (the data contract)."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="0A1430")  # eClerx navy
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _write_sheet(ws, headers: list[str], rows: list[list], widths: list[int]):
    ws.append(headers)
    for i in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=i)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _WRAP
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = _WRAP
            if cell.font is None or cell.font.size != 11:
                cell.font = Font(size=11)
    ws.freeze_panes = "A2"


def build_excel(result: dict, module_title: str) -> bytes:
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    _write_sheet(
        summary,
        ["Field", "Value"],
        [
            ["Module", module_title],
            ["Scope", result.get("scope")],
            ["Target URL", result.get("target_url")],
            ["Score", result.get("score")],
            ["Status", result.get("status")],
            ["Generated", result.get("generated_at")],
        ],
        [22, 80],
    )

    findings = wb.create_sheet("Findings")
    _write_sheet(
        findings,
        ["Signal", "Status", "Value", "Benchmark", "Source", "Evidence"],
        [
            [
                f.get("signal"), f.get("status"), f.get("value"),
                f.get("benchmark"), f.get("source"), f.get("evidence"),
            ]
            for f in result.get("findings", [])
        ],
        [26, 10, 14, 14, 48, 40],
    )

    recs = wb.create_sheet("Recommendations")
    _write_sheet(
        recs,
        ["Priority", "Layer", "Action", "How To", "Effort", "Impact"],
        [
            [
                r.get("priority"), r.get("layer"), r.get("action"),
                r.get("how_to"), r.get("effort"), r.get("impact"),
            ]
            for r in result.get("recommendations", [])
        ],
        [10, 8, 40, 60, 10, 10],
    )

    delta = wb.create_sheet("Competitor Delta")
    _write_sheet(
        delta,
        ["Competitor", "Signal", "Us", "Them", "Gap"],
        [
            [d.get("competitor"), d.get("signal"), d.get("us"), d.get("them"), d.get("gap")]
            for d in result.get("competitor_delta", [])
        ],
        [24, 24, 14, 14, 14],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
