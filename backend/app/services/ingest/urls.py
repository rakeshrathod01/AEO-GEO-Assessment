"""URL normalization + input parsing for the three ingestion modes.

(a) sitemap-URL  -> app.services.ingest.sitemap
(b) Excel upload -> parse_excel_urls
(c) paste-50     -> parse_pasted_urls
"""

from __future__ import annotations

import io
from urllib.parse import urldefrag, urlparse, urlunparse


def normalize_url(url: str) -> str | None:
    """Canonicalize a URL for dedup. Returns None if it isn't a usable http(s) URL."""
    if not url:
        return None
    url = url.strip()
    if not url:
        return None
    # Add scheme if the user pasted a bare domain. A schemeless token only counts
    # as a URL if its host looks like a domain (contains a dot) — this rejects
    # stray words like "not"/"url" that would otherwise become https://not/.
    if not urlparse(url).scheme:
        host_part = url.split("/", 1)[0]
        if "." not in host_part:
            return None
        url = "https://" + url
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        return None
    # Drop fragments, lowercase host, strip trailing slash (except root).
    url, _ = urldefrag(url)
    parsed = urlparse(url)
    netloc = parsed.netloc.lower()
    path = parsed.path or "/"
    if len(path) > 1 and path.endswith("/"):
        path = path.rstrip("/")
    return urlunparse((parsed.scheme, netloc, path, parsed.params, parsed.query, ""))


def dedupe_normalized(urls: list[str]) -> list[str]:
    """Normalize + de-duplicate while preserving first-seen order."""
    seen: set[str] = set()
    out: list[str] = []
    for raw in urls:
        n = normalize_url(raw)
        if n and n not in seen:
            seen.add(n)
            out.append(n)
    return out


def parse_pasted_urls(text: str) -> list[str]:
    """Split a pasted blob (newline/comma/space separated) into normalized URLs."""
    if not text:
        return []
    parts: list[str] = []
    for line in text.replace(",", "\n").splitlines():
        for token in line.split():
            parts.append(token)
    return dedupe_normalized(parts)


def parse_excel_urls(content: bytes, url_column: str | None = None) -> list[str]:
    """Extract URLs from an uploaded .xlsx.

    Looks for a column whose header contains "url" (case-insensitive); if none is
    found, scans every cell and keeps values that normalize to an http(s) URL.
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []

    header_cells = [str(c).strip().lower() if c is not None else "" for c in header]
    col_idx: int | None = None
    if url_column:
        target = url_column.strip().lower()
        col_idx = next((i for i, h in enumerate(header_cells) if h == target), None)
    if col_idx is None:
        col_idx = next((i for i, h in enumerate(header_cells) if "url" in h), None)

    collected: list[str] = []
    if col_idx is not None:
        # Header row's own cell might already be a URL if there were no headers.
        if normalize_url(str(header[col_idx])) is not None:
            collected.append(str(header[col_idx]))
        for row in rows:
            if col_idx < len(row) and row[col_idx]:
                collected.append(str(row[col_idx]))
    else:
        # No URL column — scan everything.
        for cell in header:
            if cell:
                collected.append(str(cell))
        for row in rows:
            for cell in row:
                if cell:
                    collected.append(str(cell))

    return dedupe_normalized(collected)
